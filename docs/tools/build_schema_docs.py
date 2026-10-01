"""Generate docs/DB-Schema-Tables.docx and docs/db-schema.drawio from schema_data.py.

    python docs/tools/build_schema_docs.py        (needs: pip install python-docx)
"""
import os
import sys
from xml.sax.saxutils import quoteattr

sys.path.insert(0, os.path.dirname(__file__))
from schema_data import ENTITIES, LAYOUT, RELATIONS  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

OUT = os.path.join(os.path.dirname(__file__), "..")
FONT = "Times New Roman"


def pretty(name):
    return " ".join(w.capitalize() for w in name.split("_"))


def set_font(run, size=13, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)


def write_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, width=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), bold=bold)
    set_cell_borders(cell)
    if width:
        cell.width = Inches(width)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    doc.styles["Normal"].font.name = FONT

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(title.add_run("Database Design: Entity Tables"), size=16, bold=True)

    widths = [0.6, 1.9, 1.5, 3.0]
    for n, (name, intro, fields) in enumerate(ENTITIES, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.keep_with_next = True
        set_font(p.add_run(f"{pretty(name)} Table: {intro}"))

        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_with_next = True
        set_font(cap.add_run(f"Table {n}. {pretty(name)} Table"), italic=True)

        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for i, h in enumerate(["No.", "Field Name", "Type", "Description"]):
            write_cell(table.rows[0].cells[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, width=widths[i])
        for idx, (field, ftype, desc) in enumerate(fields, start=1):
            row = table.add_row().cells
            write_cell(row[0], str(idx), align=WD_ALIGN_PARAGRAPH.CENTER, width=widths[0])
            write_cell(row[1], field, width=widths[1])
            write_cell(row[2], ftype, width=widths[2])
            write_cell(row[3], desc, width=widths[3])
        # repeat the header row if the table breaks across pages
        trPr = table.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trPr.append(hdr)

    path = os.path.join(OUT, "DB-Schema-Tables.docx")
    doc.save(path)
    return path


# ---------------------------------------------------------------- draw.io
ROW_H, HEAD_H, COL_W, COL_GAP, ROW_GAP = 22, 30, 270, 130, 50
HEAD_COLORS = {
    0: "#dae8fc", 1: "#d5e8d4", 2: "#fff2cc", 3: "#f8cecc", 4: "#e1d5e7",
}


def key_style(ftype):
    if "PK" in ftype:
        return "fontStyle=1"
    if "FK" in ftype:
        return "fontStyle=2"
    return ""


def build_drawio():
    by_name = {e[0]: e for e in ENTITIES}
    cells = []
    ids = {}
    next_id = [2]

    def nid():
        next_id[0] += 1
        return f"c{next_id[0]}"

    row_ids = {}  # (table, field) -> id
    for col_i, col in enumerate(LAYOUT):
        y = 20
        for name in col:
            _, _, fields = by_name[name]
            tid = nid()
            ids[name] = tid
            h = HEAD_H + ROW_H * len(fields)
            x = 20 + col_i * (COL_W + COL_GAP)
            fill = HEAD_COLORS[col_i]
            style = (
                "swimlane;fontStyle=1;childLayout=stackLayout;horizontal=1;startSize=%d;horizontalStack=0;"
                "resizeParent=1;resizeParentMax=0;resizeLast=0;collapsible=0;marginBottom=0;align=center;"
                "fontSize=13;fontFamily=Times New Roman;fillColor=%s;strokeColor=#000000;swimlaneFillColor=#ffffff;"
                % (HEAD_H, fill)
            )
            cells.append(
                f'<mxCell id="{tid}" value={quoteattr(name)} style="{style}" vertex="1" parent="1">'
                f'<mxGeometry x="{x}" y="{y}" width="{COL_W}" height="{h}" as="geometry"/></mxCell>'
            )
            for fi, (field, ftype, _) in enumerate(fields):
                rid = nid()
                row_ids[(name, field)] = rid
                label = f"{field} : {ftype}"
                rstyle = (
                    "text;strokeColor=none;fillColor=none;align=left;verticalAlign=middle;spacingLeft=6;"
                    "spacingRight=4;overflow=hidden;rotatable=0;points=[[0,0.5],[1,0.5]];portConstraint=eastwest;"
                    "fontSize=11;fontFamily=Times New Roman;" + key_style(ftype)
                )
                cells.append(
                    f'<mxCell id="{rid}" value={quoteattr(label)} style="{rstyle}" vertex="1" parent="{tid}">'
                    f'<mxGeometry y="{HEAD_H + fi * ROW_H}" width="{COL_W}" height="{ROW_H}" as="geometry"/></mxCell>'
                )
            y += h + ROW_GAP

    for child, cfield, parent, pfield, kind in RELATIONS:
        eid = nid()
        start = "ERmandOne" if kind == "11" else "ERmany"
        style = (
            "edgeStyle=entityRelationEdgeStyle;fontSize=11;html=1;endArrow=ERmandOne;startArrow=%s;"
            "endFill=0;startFill=0;strokeColor=#333333;rounded=0;" % ("ERmandOne" if kind == "11" else "ERmany")
        )
        src = row_ids[(child, cfield)]
        dst = row_ids[(parent, pfield)]
        cells.append(
            f'<mxCell id="{eid}" style="{style}" edge="1" parent="1" source="{src}" target="{dst}">'
            f'<mxGeometry relative="1" as="geometry"/></mxCell>'
        )

    total_w = 40 + len(LAYOUT) * (COL_W + COL_GAP)
    xml = (
        '<mxfile host="app.diagrams.net" agent="build_schema_docs.py" version="24.0.0">\n'
        '  <diagram id="officehours-erd" name="OfficeHours ERD">\n'
        f'    <mxGraphModel dx="1600" dy="1000" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" '
        f'arrows="1" fold="1" page="1" pageScale="1" pageWidth="{total_w}" pageHeight="1800" math="0" shadow="0">\n'
        "      <root>\n"
        '        <mxCell id="0"/>\n'
        '        <mxCell id="1" parent="0"/>\n'
        + "\n".join("        " + c for c in cells)
        + "\n      </root>\n    </mxGraphModel>\n  </diagram>\n</mxfile>\n"
    )
    path = os.path.join(OUT, "db-schema.drawio")
    with open(path, "w", encoding="utf-8") as f:
        f.write(xml)
    return path


if __name__ == "__main__":
    print(build_docx())
    print(build_drawio())
